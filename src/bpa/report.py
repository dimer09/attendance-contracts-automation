import datetime as dt
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bpa.models import RuleViolation, Severity
from bpa.validate import RejectedRow

RUN_ID_PATTERN = re.compile(r"[A-Za-z0-9_-]+")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SEVERITY_FILLS = {
    Severity.BLOCKING: PatternFill("solid", fgColor="F8CBAD"),  
    Severity.WARNING: PatternFill("solid", fgColor="FFE699"),   
}
DATE_FORMAT = "yyyy-mm-dd"

EXCEPTION_HEADERS = ["Rule", "Severity", "Employee ID", "Name", "Client", "Date", "Hours", "Message"]
EXCEPTION_WIDTHS = [8, 12, 14, 20, 18, 12, 8, 70]

REJECTED_HEADERS = ["File row", "Employee ID", "Name", "Client", "Date", "Hours", "Reason"]
REJECTED_WIDTHS = [10, 14, 20, 18, 14, 10, 70]
RAW_KEYS = ["employee_id", "name", "client", "date", "hours"]


def _set(sheet: Worksheet, row: int, column: int, value, number_format: str | None = None) -> Cell:
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    if number_format:
        cell.number_format = number_format
    return cell


def _write_header(sheet: Worksheet, headers: list[str], widths: list[int]) -> None:
    for column, (title, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = _set(sheet, 1, column, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"  


def result_label(blocking: int, warnings: int, rejected_count: int) -> str:
    if blocking or rejected_count:
        return "ACTION REQUIRED"
    if warnings:
        return "REVIEW RECOMMENDED"
    return "ALL CLEAR"


def _write_summary(
    sheet: Worksheet,
    *,
    run_id: str,
    source_name: str,
    generated_at: dt.datetime,
    valid_count: int,
    violations: list[RuleViolation],
    rejected: list[RejectedRow],
) -> None:
    severity_counts = Counter(violation.severity for violation in violations)
    blocking = severity_counts[Severity.BLOCKING]
    warnings = severity_counts[Severity.WARNING]

    entries = [
        ("Run ID", run_id),
        ("Source file", source_name),
        ("Generated at", generated_at.isoformat(timespec="seconds")),
        None,  # blank line
        ("Lines read", valid_count + len(rejected)),
        ("Valid lines", valid_count),
        ("Rejected lines", len(rejected)),
        ("Exceptions", len(violations)), 
        ("Blocking exceptions", blocking),
        ("Warnings", warnings),
        None,
        ("Result", result_label(blocking, warnings, len(rejected))),
    ]

    row = 1
    for entry in entries:
        if entry is not None:
            label, value = entry
            _set(sheet, row, 1, label).font = Font(bold=True)
            _set(sheet, row, 2, value)
        row += 1

    row += 1  
    _set(sheet, row, 1, "Exceptions by rule").font = Font(bold=True)
    per_rule = Counter(violation.rule_code for violation in violations)
    for code in sorted(per_rule):
        row += 1
        _set(sheet, row, 1, code)
        _set(sheet, row, 2, per_rule[code])

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 40


def _write_exceptions(sheet: Worksheet, violations: list[RuleViolation]) -> None:
    _write_header(sheet, EXCEPTION_HEADERS, EXCEPTION_WIDTHS)
    for row, violation in enumerate(violations, start=2):
        record = violation.record
        values = [
            violation.rule_code,
            violation.severity.value,
            record.employee_id,
            record.name,
            record.client,
            record.date,  
            record.hours,  
            violation.message,
        ]
        for column, value in enumerate(values, start=1):
            cell = _set(sheet, row, column, value, DATE_FORMAT if isinstance(value, dt.date) else None)
            cell.fill = SEVERITY_FILLS[violation.severity]
    sheet.auto_filter.ref = sheet.dimensions  


def _write_rejected(sheet: Worksheet, rejected: list[RejectedRow]) -> None:
    _write_header(sheet, REJECTED_HEADERS, REJECTED_WIDTHS)
    for row, item in enumerate(rejected, start=2):
        values = [item.row_number, *(item.raw.get(key, "") for key in RAW_KEYS), item.reason]
        for column, value in enumerate(values, start=1):
            _set(sheet, row, column, value)
    sheet.auto_filter.ref = sheet.dimensions


def generate_report(
    output_dir: Path,
    *,
    run_id: str,
    source_name: str,
    valid_count: int,
    violations: list[RuleViolation],
    rejected: list[RejectedRow],
    generated_at: dt.datetime | None = None,
) -> Path:
    """Write the Excel report and return its path."""
    if not RUN_ID_PATTERN.fullmatch(run_id):
        raise ValueError(f"Unsafe run id: {run_id!r}")
    generated_at = generated_at or dt.datetime.now()

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(
        summary,
        run_id=run_id,
        source_name=source_name,
        generated_at=generated_at,
        valid_count=valid_count,
        violations=violations,
        rejected=rejected,
    )
    _write_exceptions(workbook.create_sheet("Exceptions"), violations)
    _write_rejected(workbook.create_sheet("Rejected rows"), rejected)

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"report_{run_id}.xlsx"  
    workbook.save(path)
    return path