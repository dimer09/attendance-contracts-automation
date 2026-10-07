import json
import os
import socket
import subprocess
import sys
import threading
import time
from pathlib import Path

import pandas as pd
import pytest
import uvicorn
from openpyxl import load_workbook

from mock_api import main as mock_main

pytestmark = pytest.mark.e2e 

ATTENDANCE = Path(__file__).parent.parent / "data" / "sample" / "attendance.csv"
E2E_KEY = "e2e-secret-key-4f9a"


@pytest.fixture(scope="module")
def api_url():
    """The mock API of step 5, really listening on a free local port."""

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(mock_main, "API_KEY", E2E_KEY)

        config = uvicorn.Config(
            mock_main.app, host="127.0.0.1", port=0, log_config=None, log_level="warning"
        )
        server = uvicorn.Server(config)
        thread = threading.Thread(target=server.run, daemon=True)
        thread.start()
        try:
            deadline = time.monotonic() + 10
            while not server.started:  
                if time.monotonic() > deadline or not thread.is_alive():
                    pytest.fail("The mock API did not start")
                time.sleep(0.05)
            port = server.servers[0].sockets[0].getsockname()[1]
            yield f"http://127.0.0.1:{port}"
        finally:
            server.should_exit = True
            thread.join(timeout=5)


@pytest.fixture
def unused_url():
    """URL of a local port where nothing is listening."""
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    return f"http://127.0.0.1:{port}"  


def run_bpa(tmp_path, input_path, *, url, key=E2E_KEY):

    env = {k: v for k, v in os.environ.items() if not k.startswith("BPA_")}
    env.update({
        "BPA_API_URL": url,
        "BPA_API_KEY": key,
        "NO_PROXY": "127.0.0.1,localhost",
        "no_proxy": "127.0.0.1,localhost",
        "PYTHONIOENCODING": "utf-8",
    })
    return subprocess.run(
        [sys.executable, "-m", "bpa", "--input", str(input_path)],  
        cwd=tmp_path, 
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
    )


def read_log(tmp_path):
    lines = (tmp_path / "logs" / "bpa.log").read_text(encoding="utf-8").splitlines()
    return [json.loads(line) for line in lines]


def stdout_value(result, label):
  
    for line in result.stdout.splitlines():
        if line.startswith(f"{label}: "):
            return line[len(label) + 2:]
    raise AssertionError(f"'{label}' not found in output:\n{result.stdout}")


def test_full_run_on_sample_data(api_url, tmp_path):
    result = run_bpa(tmp_path, ATTENDANCE, url=api_url)

    assert result.returncode == 0, result.stderr
    assert "Exceptions: 6" in result.stdout
    assert "Result: ACTION REQUIRED" in result.stdout

    assert "INFO" not in result.stdout
    assert "INFO" in result.stderr
   
    run_id = stdout_value(result, "Run ID")
    report = tmp_path / "reports" / f"report_{run_id}.xlsx"
    assert load_workbook(report).sheetnames == ["Summary", "Exceptions", "Rejected rows"]
    assert {entry["run_id"] for entry in read_log(tmp_path)} == {run_id}


def test_excel_input_gives_the_same_result_as_csv(api_url, tmp_path):
    excel = tmp_path / "attendance.xlsx"
    pd.read_csv(ATTENDANCE, dtype=str, keep_default_na=False).to_excel(excel, index=False)

    result = run_bpa(tmp_path, excel, url=api_url)

    assert result.returncode == 0, result.stderr
    assert "Lines read: 12 (valid: 10, rejected: 2)" in result.stdout
    assert "Exceptions: 6" in result.stdout


def test_api_down_exits_2_after_retries(unused_url, tmp_path):
    result = run_bpa(tmp_path, ATTENDANCE, url=unused_url) 

    assert result.returncode == 2, result.stderr
    assert "unavailable after 3 attempts" in result.stderr
    assert result.stderr.count("Retrying") == 2  
    assert result.stdout == ""                  
    assert not (tmp_path / "reports").exists()  


def test_wrong_key_exits_2_without_retry(api_url, tmp_path):
    result = run_bpa(tmp_path, ATTENDANCE, url=api_url, key="wrong-key")

    assert result.returncode == 2, result.stderr
    assert "status 401" in result.stderr
    assert "Retrying" not in result.stderr  


def test_missing_input_exits_1_before_any_api_call(unused_url, tmp_path):
    result = run_bpa(tmp_path, tmp_path / "nope.csv", url=unused_url)

    assert result.returncode == 1, result.stderr
    assert "File not found" in result.stderr
  
    assert "Retrying" not in result.stderr


@pytest.mark.parametrize(
    "file_name, content, message",
    [
        ("empty.csv", "", "Error reading file"),
        ("wrong_columns.csv", "id,name\nM001,Mbuyi\n", "Missing required columns:"),
        ("notes.txt", "employee_id,name,client,date,hours\n", "Unsupported file extension"),
    ],
)
def test_unusable_input_exits_1(unused_url, tmp_path, file_name, content, message):
    path = tmp_path / file_name
    path.write_text(content)

    result = run_bpa(tmp_path, path, url=unused_url)

    assert result.returncode == 1, result.stderr
    assert message in result.stderr


def test_header_only_file_is_all_clear(api_url, tmp_path):
    path = tmp_path / "empty_week.csv"
    path.write_text("employee_id,name,client,date,hours\n")

    result = run_bpa(tmp_path, path, url=api_url)

    assert result.returncode == 0, result.stderr
    assert "Lines read: 0" in result.stdout
    assert "Result: ALL CLEAR" in result.stdout


def test_only_invalid_rows_still_requires_action(api_url, tmp_path):
    path = tmp_path / "broken.csv"
    path.write_text(
        "employee_id,name,client,date,hours\n"
        "M001,Mbuyi,Acme,not-a-date,8\n"
        "M001,Mbuyi,Acme,2026-10-05,-3\n"
    )

    result = run_bpa(tmp_path, path, url=api_url)

    assert result.returncode == 0, result.stderr
    assert "Exceptions: 0" in result.stdout
    assert "rejected: 2" in result.stdout
    assert "Result: ACTION REQUIRED" in result.stdout 


@pytest.mark.parametrize("key", [E2E_KEY, "wrong-key-7c1d9b"])
def test_api_key_never_appears_in_output_or_logs(api_url, tmp_path, key):
    result = run_bpa(tmp_path, ATTENDANCE, url=api_url, key=key)

    log_text = (tmp_path / "logs" / "bpa.log").read_text(encoding="utf-8")
    assert log_text.strip() 
    assert key not in result.stdout
    assert key not in result.stderr
    assert key not in log_text


def test_two_runs_keep_separate_reports_and_logs(api_url, tmp_path):
    for _ in range(2):
        result = run_bpa(tmp_path, ATTENDANCE, url=api_url)
        assert result.returncode == 0, result.stderr

    assert len(list((tmp_path / "reports").glob("report_*.xlsx"))) == 2  
    assert len({entry["run_id"] for entry in read_log(tmp_path)}) == 2