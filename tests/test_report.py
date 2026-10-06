import datetime as dt
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bpa.compare import match_records_to_contracts
from bpa.extract.files import read_attendance_file
from bpa.models import Contract, RuleViolation, Severity
from bpa.report import generate_report
from bpa.rules import apply_rules
from bpa.validate import RejectedRow, validate_attendance_rows
from tests.factories import make_record

SAMPLE_DIR = Path(__file__).parent.parent / "data" / "sample"
GENERATED_AT = dt.datetime(2026, 10, 6, 14, 30, 0)
RUN_ID = "20261006-143000"


def make_violation(rule_code="R1", severity=Severity.BLOCKING, **record_overrides):
    record = make_record(**record_overrides)
    return RuleViolation(rule_code, severity, record, f"{rule_code} problem")


def make_rejected(row_number=3):
    raw = {
        "employee_id": "M001",
        "name": "Mbuyi",
        "client": "Acme",
        "date": "2026-10-07",
        "hours": "-3",
    }
    return RejectedRow(row_number, raw, "hours: Input should be greater than 0")


def build(tmp_path, *, valid_count=0, violations=(), rejected=()):
    return generate_report(
        tmp_path,
        run_id=RUN_ID,
        source_name="attendance.csv",
        valid_count=valid_count,
        violations=list(violations),
        rejected=list(rejected),
        generated_at=GENERATED_AT,
    )


def summary_of(path):
    """The Summary sheet as {label: value}."""
    sheet = load_workbook(path)["Summary"]
    return {row[0]: row[1] for row in sheet.iter_rows(values_only=True) if row[0]}


def test_creates_one_file_with_three_sheets(tmp_path):
    path = build(tmp_path)

    assert path.name == f"report_{RUN_ID}.xlsx"
    assert load_workbook(path).sheetnames == ["Summary", "Exceptions", "Rejected rows"]


def test_summary_shows_counts_and_rule_breakdown(tmp_path):
    violations = [
        make_violation("R1"),
        make_violation("R1", employee_id="M002"),
        make_violation("R3", Severity.WARNING),
    ]

    summary = summary_of(build(tmp_path, valid_count=5, violations=violations, rejected=[make_rejected()]))

    assert summary["Run ID"] == RUN_ID
    assert summary["Source file"] == "attendance.csv"
    assert summary["Generated at"] == "2026-10-06T14:30:00"
    assert summary["Lines read"] == 6      
    assert summary["Valid lines"] == 5
    assert summary["Rejected lines"] == 1
    assert summary["Exceptions"] == 3
    assert summary["Blocking exceptions"] == 2
    assert summary["Warnings"] == 1
    assert summary["R1"] == 2
    assert summary["R3"] == 1
    assert summary["Result"] == "ACTION REQUIRED"


def test_exceptions_sheet_has_values_and_severity_colors(tmp_path):
    violations = [make_violation("R1"), make_violation("R3", Severity.WARNING)]

    sheet = load_workbook(build(tmp_path, violations=violations))["Exceptions"]

    headers = [cell.value for cell in sheet[1]]
    assert headers == ["Rule", "Severity", "Employee ID", "Name", "Client", "Date", "Hours", "Message"]
    first = [cell.value for cell in sheet[2]]
    assert first == ["R1", "blocking", "M001", "Mbuyi", "Acme", dt.datetime(2026, 10, 5), 8, "R1 problem"]
    assert sheet["A2"].fill.fgColor.rgb.endswith("F8CBAD") 
    assert sheet["A3"].fill.fgColor.rgb.endswith("FFE699") 
    assert sheet.freeze_panes == "A2"


def test_rejected_sheet_keeps_original_values_and_reason(tmp_path):
    sheet = load_workbook(build(tmp_path, rejected=[make_rejected(row_number=12)]))["Rejected rows"]

    row = [cell.value for cell in sheet[2]]
    assert row == [12, "M001", "Mbuyi", "Acme", "2026-10-07", "-3", "hours: Input should be greater than 0"]


def test_text_starting_with_equals_stays_text(tmp_path):

    path = build(tmp_path, violations=[make_violation(name="=1+1")])

    cell = load_workbook(path)["Exceptions"]["D2"]

    assert cell.value == "=1+1"
    assert cell.data_type == "s"  


def test_empty_report_has_headers_only(tmp_path):
    workbook = load_workbook(build(tmp_path))

    assert workbook["Exceptions"].max_row == 1
    assert workbook["Rejected rows"].max_row == 1


@pytest.mark.parametrize("bad_run_id", ["../evil", "a/b", ""])
def test_unsafe_run_id_is_rejected(tmp_path, bad_run_id):
    with pytest.raises(ValueError, match="Unsafe run id"):
        generate_report(
            tmp_path, run_id=bad_run_id, source_name="x.csv",
            valid_count=0, violations=[], rejected=[],
        )


def test_missing_output_directory_is_created(tmp_path):
    target = tmp_path / "out" / "reports"

    path = generate_report(
        target, run_id=RUN_ID, source_name="x.csv",
        valid_count=0, violations=[], rejected=[],
    )

    assert path.parent == target
    assert path.exists()


@pytest.mark.parametrize(
    "violations, rejected, expected",
    [
        ([], [], "ALL CLEAR"),
        ([make_violation("R3", Severity.WARNING)], [], "REVIEW RECOMMENDED"),
        ([make_violation("R1")], [], "ACTION REQUIRED"),
        ([], [make_rejected()], "ACTION REQUIRED"),  
    ],
)
def test_result_reflects_what_was_found(tmp_path, violations, rejected, expected):
    path = build(tmp_path, violations=violations, rejected=rejected)

    assert summary_of(path)["Result"] == expected


def test_sample_data_end_to_end(tmp_path):

    result = validate_attendance_rows(read_attendance_file(SAMPLE_DIR / "attendance.csv"))
    raw_contracts = json.loads((SAMPLE_DIR / "contracts.json").read_text())
    contracts = [Contract(**item) for item in raw_contracts]
    violations = apply_rules(match_records_to_contracts(result.valid, contracts))

    path = generate_report(
        tmp_path, run_id="sample", source_name="attendance.csv",
        valid_count=len(result.valid), violations=violations,
        rejected=result.rejected, generated_at=GENERATED_AT,
    )

    summary = summary_of(path)
    assert summary["Lines read"] == 12
    assert summary["Valid lines"] == 10
    assert summary["Rejected lines"] == 2
    assert summary["Exceptions"] == 6
    assert summary["Blocking exceptions"] == 5
    assert summary["Warnings"] == 1
    assert summary["R4"] == 2
    assert summary["Result"] == "ACTION REQUIRED"