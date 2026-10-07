import json
import logging
from pathlib import Path

import pytest
import responses
from openpyxl import load_workbook

from bpa.config import RulesConfig, RuleSetting
from bpa.extract import api_client
from bpa.extract.api_client import ApiError
from bpa.extract.files import FileReaderError
from bpa.pipeline import run_pipeline

SAMPLE_DIR = Path(__file__).parent.parent / "data" / "sample"
ATTENDANCE = SAMPLE_DIR / "attendance.csv"
BASE_URL = "http://api.test"
CONTRACTS_URL = f"{BASE_URL}/contracts"
CONTRACTS = json.loads((SAMPLE_DIR / "contracts.json").read_text())


@pytest.fixture(autouse=True)
def no_real_waiting(monkeypatch):
    monkeypatch.setattr(api_client.time, "sleep", lambda seconds: None)


def serve_contracts():
    responses.add(responses.GET, CONTRACTS_URL, json=CONTRACTS, status=200)


def run(tmp_path, input_path=ATTENDANCE, run_id="test-run", **kwargs):
    return run_pipeline(
        input_path,
        base_url=BASE_URL,
        api_key="key",
        output_dir=tmp_path / "reports",
        run_id=run_id,
        **kwargs,
    )


@responses.activate
def test_sample_run_gives_the_expected_counts(tmp_path):
    serve_contracts()

    result = run(tmp_path)

    assert result.lines_read == 12
    assert result.valid_count == 10
    assert len(result.rejected) == 2
    assert len(result.violations) == 6
    assert result.blocking_count == 5
    assert result.warning_count == 1
    assert result.verdict == "ACTION REQUIRED"


@responses.activate
def test_report_file_belongs_to_the_run(tmp_path):
    serve_contracts()

    result = run(tmp_path, run_id="run-42")

    assert result.report_path.name == "report_run-42.xlsx"
    sheet = load_workbook(result.report_path)["Summary"]
    values = {row[0]: row[1] for row in sheet.iter_rows(values_only=True) if row[0]}
    assert values["Source file"] == "attendance.csv"
    assert values["Result"] == result.verdict 


@responses.activate
def test_missing_input_file_fails_before_calling_the_api(tmp_path):
    with pytest.raises(FileReaderError):
        run(tmp_path, input_path=tmp_path / "nope.csv")

    assert len(responses.calls) == 0  


@responses.activate
def test_api_failure_leaves_no_report(tmp_path):
    for _ in range(3):
        responses.add(responses.GET, CONTRACTS_URL, status=503)

    with pytest.raises(ApiError):
        run(tmp_path)

    assert not (tmp_path / "reports").exists() 


@responses.activate
def test_same_inputs_give_same_exceptions_in_separate_reports(tmp_path):
    serve_contracts()

    first = run(tmp_path, run_id="run-a")
    second = run(tmp_path, run_id="run-b")

    assert first.violations == second.violations
    assert first.report_path != second.report_path
    assert first.report_path.exists() and second.report_path.exists()


@responses.activate
def test_rules_config_is_applied(tmp_path):
    serve_contracts()
    config = RulesConfig(settings={"R4": RuleSetting(enabled=False)})

    result = run(tmp_path, rules_config=config)

    assert len(result.violations) == 4  
    assert all(v.rule_code != "R4" for v in result.violations)


@responses.activate
def test_rejected_rows_are_logged_without_personal_names(tmp_path, caplog):
    caplog.set_level(logging.INFO)
    serve_contracts()

    run(tmp_path)

    assert "Rejected row 11" in caplog.text
    assert "Rejected row 12" in caplog.text
    for name in ("Mbuyi", "Kalala", "Tshala", "Ilunga"):
        assert name not in caplog.text  

@responses.activate
def test_header_only_file_gives_all_clear(tmp_path):
    serve_contracts()
    empty = tmp_path / "empty.csv"
    empty.write_text("employee_id,name,client,date,hours\n")

    result = run(tmp_path, input_path=empty)

    assert result.lines_read == 0
    assert result.verdict == "ALL CLEAR"